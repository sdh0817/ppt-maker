#!/usr/bin/env python3
"""Generate a research-publication PPT from a fixed list."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PAPERS = [
    ("2007-08", "제어 가능한 전송 영점을 이용한 광대역 차단 특성을 갖는 저역 통과 필터"),
    ("2007-08", "단락 개방 Calibration 방법을 이용한 MIM 커패시터의 기생 소자 값 추출"),
    ("2009-08", "LTCC 적층 필터를 위한 기생 성분 해석 및 필터 설계"),
    ("2010-08", "향상된 보정 기술을 이용한 수동 소자의 동적 모델링과 고주파 필터 설계에의 응용",
     "박사 졸업 논문"),
    ("2017-12", "공통모드 초크의 간단한 고주파 모델링 기법"),
    ("2022-04", "향상된 등가 모델링 기법을 이용한 소형 대역통과 필터 설계"),
    ("2025-05", "위상 천이기와 대역통과 필터 복합 기능 구현으로 슬림 TV의 Wi-Fi·Bluetooth 모듈 소형화 및 양산화",
     "해동 산업체 우수논문상 수상"),
]

BG_DARK = RGBColor(0x0A, 0x0A, 0x0A)
FG_LIGHT = RGBColor(0xEE, 0xEE, 0xEE)
FG_DIM = RGBColor(0xA0, 0xA0, 0xA0)
ACCENT = RGBColor(0xFF, 0xD0, 0x00)
FONT = "맑은 고딕"


def set_solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, x, y, w, h, text, *, size=18, color=FG_LIGHT,
             bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def main():
    out = Path(__file__).resolve().parent.parent / "projects" / "research-portfolio"
    out.mkdir(parents=True, exist_ok=True)
    out_file = out / "research_papers.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    # --- Slide 1: Title ---
    s = prs.slides.add_slide(blank)
    set_solid_bg(s, BG_DARK)
    add_text(s, Inches(0.8), Inches(2.3), Inches(12), Inches(1.4),
             "연구 논문 리스트", size=54, color=FG_LIGHT, bold=True)
    add_text(s, Inches(0.8), Inches(3.6), Inches(12), Inches(0.6),
             "2007 – 2025 · 7편", size=22, color=ACCENT)
    add_text(s, Inches(0.8), Inches(6.7), Inches(12), Inches(0.4),
             "고주파 회로 · 수동 소자 모델링 · 필터 설계",
             size=14, color=FG_DIM)

    # --- Slide 2: Timeline overview ---
    s = prs.slides.add_slide(blank)
    set_solid_bg(s, BG_DARK)
    add_text(s, Inches(0.8), Inches(0.4), Inches(12), Inches(0.8),
             "논문 발표 타임라인", size=32, color=FG_LIGHT, bold=True)

    y0 = Inches(1.6)
    row_h = Inches(0.75)
    for i, p in enumerate(PAPERS):
        date, title = p[0], p[1]
        badge = p[2] if len(p) > 2 else None
        y = y0 + row_h * i
        add_text(s, Inches(0.9), y, Inches(1.5), Inches(0.6),
                 date, size=18, color=ACCENT, bold=True)
        add_text(s, Inches(2.5), y, Inches(9.5), Inches(0.6),
                 title, size=15, color=FG_LIGHT)
        if badge:
            add_text(s, Inches(2.5), y + Inches(0.35), Inches(9.5), Inches(0.3),
                     f"— {badge}", size=11, color=ACCENT)

    # --- Slides 3+: One per paper ---
    for i, p in enumerate(PAPERS, start=1):
        date, title = p[0], p[1]
        badge = p[2] if len(p) > 2 else None
        s = prs.slides.add_slide(blank)
        set_solid_bg(s, BG_DARK)

        add_text(s, Inches(0.8), Inches(0.5), Inches(2.0), Inches(0.6),
                 f"#{i}", size=48, color=ACCENT, bold=True)
        add_text(s, Inches(0.8), Inches(1.6), Inches(6), Inches(0.5),
                 date, size=22, color=FG_DIM, bold=True)
        add_text(s, Inches(0.8), Inches(2.3), Inches(12), Inches(3.5),
                 title, size=32, color=FG_LIGHT, bold=True)
        if badge:
            add_text(s, Inches(0.8), Inches(5.9), Inches(12), Inches(0.7),
                     f"★  {badge}", size=20, color=ACCENT, bold=True)

    prs.save(out_file)
    print(f"[done] {out_file}  ({out_file.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
